"""Renderer tests for the 3-section report (markdown / html / docx).

Board variant only. Verifies the renderers are deterministic, emit the three
README sections, produce one file per format, and leave no unresolved
`{placeholder}` strings.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import pytest
from docx import Document

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))

from src.reporting.industry_analysis_report import build_report  # noqa: E402
from src.reporting.render_docx import write_docx_variants  # noqa: E402
from src.reporting.render_html import render as render_html  # noqa: E402
from src.reporting.render_markdown import (  # noqa: E402
    render as render_markdown,
    write_markdown_variants,
)
from src.reporting.render_html import write_html_variants  # noqa: E402


@pytest.fixture(scope="module")
def report():
    return build_report()


def test_markdown_output_deterministic(report):
    assert render_markdown(report, "board") == render_markdown(report, "board")


def test_board_markdown_has_the_three_sections(report):
    md = render_markdown(report, "board")
    assert "## 1. Macro conditions for credit assessment & risk management" in md
    assert "## 2. Macro drivers for stress testing - per product and per industry" in md
    assert "## 3. How the industry credit-risk score is calculated" in md


def test_html_has_no_broken_placeholders(report):
    rendered = render_html(report, "board")
    assert rendered.startswith("<!DOCTYPE html>")
    assert "<style>" in rendered
    assert not re.findall(r"\{[a-z_][a-z0-9_]*\}", rendered)


def test_each_format_produces_one_board_file(tmp_path, report):
    base = tmp_path / "Industry_Analysis_Test"
    for writer in (write_markdown_variants, write_html_variants, write_docx_variants):
        result = writer(report, str(base))
        assert set(result) == {"board"}
        for path in result.values():
            assert Path(path).exists()
            assert Path(path).stat().st_size > 1024


def test_docx_contains_section_titles_and_no_appendix(tmp_path, report):
    base = tmp_path / "Industry_Analysis_DocxTest"
    paths = write_docx_variants(report, str(base))
    doc = Document(paths["board"])
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "1. Macro conditions for credit assessment & risk management" in text
    assert "3. How the industry credit-risk score is calculated" in text
    assert "Appendix A" not in text
